"""Zero-provider, zero-SMTP tests for the one-click Salesoorja operator."""

from __future__ import annotations

import threading
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from services.rediff_sender_adapter import RediffAdapterConfig, RediffSenderAdapter
from services.rediff_transport_bridge import TEST_TRANSPORT_AUTHORIZATION, RediffTransportBridge
from services.salesoorja_operator import REPORT_SHEETS, SalesoorjaOperator


NOW = datetime(2026, 9, 14, 10, 30, tzinfo=timezone.utc)


def _settings(**overrides):
    values = {
        "SALESOORJA_MODE": "TEST",
        "REAL_OUTREACH_ENABLED": False,
        "OUTBOUND_TEST_MODE": True,
        "REDIFF_SENDER_ENABLED": False,
        "REDIFF_TEST_MODE": True,
        "SERPER_API_KEY": "",
        "APOLLO_API_KEY": "",
        "BRIGHTDATA_API_TOKEN": "",
        "BRIGHTDATA_LINKEDIN_PROFILE_DATASET_ID": "",
        "HIVE_API_KEY": "",
        "GEMINI_API_KEY": "",
        "GOOGLE_API_KEY": "",
        "DAILY_SEND_TARGET": 150,
        "DAILY_SEND_MAX": 250,
        "MAX_SERPER_CALLS_PER_DAY": 1500,
        "SALESOORJA_START_TIME": "09:00",
        "SALESOORJA_END_TIME": "18:00",
        "SALESOORJA_CYCLE_INTERVAL_SECONDS": 1,
        "SALESOORJA_INBOX_INTERVAL_SECONDS": 1,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def _operator(tmp_path: Path, db_factory=None, transport_bridge=None, **settings_overrides) -> SalesoorjaOperator:
    rediff = RediffSenderAdapter(
        config=RediffAdapterConfig(
            enabled=False,
            test_mode=True,
            system_path=tmp_path / "missing-rediff",
            handoff_dir=tmp_path / "handoff",
            cc_addresses=("Bablu@oorjatechnical.org", "piyushk@oorjatechnical.com"),
            duplicate_window_days=14,
        ),
        now=lambda: NOW,
    )
    return SalesoorjaOperator(
        settings_obj=_settings(**settings_overrides),
        state_path=tmp_path / "runtime" / "operator_state.json",
        report_dir=tmp_path / "reports",
        now=lambda: NOW,
        db_factory=db_factory,
        rediff_adapter=rediff,
        transport_bridge=transport_bridge,
    )


def test_complete_synthetic_run_generates_expected_report_without_smtp(tmp_path, monkeypatch):
    smtp_calls = []
    monkeypatch.setattr("smtplib.SMTP", lambda *args, **kwargs: smtp_calls.append((args, kwargs)))
    monkeypatch.setattr("smtplib.SMTP_SSL", lambda *args, **kwargs: smtp_calls.append((args, kwargs)))
    operator = _operator(tmp_path, SALESOORJA_MODE="SMOKE")

    result = operator.start_run(background=False)
    status = result["status"]

    assert result["started"] is True
    assert status["status"] == "STOPPED"
    assert status["stop_reason"] == "SMOKE_COMPLETE"
    assert status["counters"]["companies_researched"] == 1
    assert status["counters"]["qualified_opportunities"] == 1
    assert status["counters"]["people_verified"] == 1
    assert status["counters"]["contacts_enriched"] == 1
    assert status["counters"]["emails_sent"] == 0
    assert status["provider_usage"]["Rediff"] == 1
    assert status["provider_usage"]["Serper"] == 0
    assert smtp_calls == []

    report = Path(status["report_path"])
    assert report.is_file()
    workbook = load_workbook(report, read_only=True)
    assert workbook.sheetnames == ["SUMMARY", *REPORT_SHEETS]
    assert status["final_report_email"]["status"] == "DRY_RUN_READY"
    assert status["final_report_email"]["smtp_sent"] is False
    assert status["final_report_email"]["attachment"] == str(report)


def test_restart_resets_run_scoped_state_and_preserves_history(tmp_path):
    first = _operator(tmp_path, SALESOORJA_MODE="SMOKE")
    first.start_run(background=False)
    first_status = first.get_status()

    resumed = _operator(tmp_path, SALESOORJA_MODE="SMOKE")
    second_result = resumed.start_run(background=False)
    second_status = second_result["status"]

    assert second_status["run_id"] != first_status["run_id"]
    assert second_status["counters"]["companies_researched"] == 1
    assert second_status["historical_counters"]["companies_researched"] == 1
    assert second_status["processed_accounts"] == ["synthetic:operator-control"]


def test_start_state_resets_all_transient_run_fields(tmp_path):
    operator = _operator(tmp_path)
    operator._state.update(
        current_company="Jabil",
        last_action="Stale discovery",
        report_path="stale.xlsx",
        final_report_email={"status": "STALE"},
        processed_accounts=["stale-company"],
    )
    operator._state["provider_usage"]["Serper"] = 9
    operator._state["counters"]["companies_researched"] = 3
    operator._state["records"]["HOLD"].append({"company": "Jabil"})

    state = operator._new_state(preserve_history=True)

    assert state["current_company"] is None
    assert state["last_action"] == "Ready"
    assert state["report_path"] is None
    assert state["final_report_email"] is None
    assert state["processed_accounts"] == []
    assert state["provider_usage"]["Serper"] == 0
    assert state["counters"]["companies_researched"] == 0
    assert state["records"]["HOLD"] == []


def test_test_mode_repeats_real_discovery_until_manual_stop(tmp_path):
    operator = _operator(tmp_path, db_factory=lambda: SimpleNamespace(close=lambda: None))
    cycles = []
    operator._poll_inbox = lambda: None
    operator._process_due_followups = lambda db: 0

    def run_cycle():
        cycles.append(len(cycles) + 1)
        if len(cycles) == 2:
            operator._set_stop_signal()
        return True

    operator._run_discovery_cycle = run_cycle
    operator._state["status"] = "RUNNING"

    operator._run_test_mode()

    assert cycles == [1, 2]
    assert operator.get_status()["stop_reason"] == "MANUAL_STOP"


def test_production_mode_retries_discovery_and_processes_followups_each_cycle(tmp_path):
    database = SimpleNamespace(close=lambda: None)
    operator = _operator(
        tmp_path,
        db_factory=lambda: database,
        SALESOORJA_MODE="PRODUCTION",
        REAL_OUTREACH_ENABLED=True,
        OUTBOUND_TEST_MODE=False,
    )
    cycles = []
    followup_cycles = []
    waits = []
    operator._poll_inbox = lambda: None

    def run_cycle():
        cycles.append(len(cycles) + 1)
        if len(cycles) == 1:
            raise RuntimeError("temporary provider failure")
        return False

    def wait_for_next_cycle(seconds):
        waits.append(seconds)
        if len(waits) == 2:
            operator._set_stop_signal()
            return True
        return False

    operator._run_discovery_cycle = run_cycle
    operator._process_due_followups = lambda db: followup_cycles.append(db) or 0
    operator._interruptible_wait = wait_for_next_cycle
    operator._state["status"] = "RUNNING"

    operator._run_production_loop()

    assert cycles == [1, 2]
    assert followup_cycles == [database, database]
    assert operator.get_status()["counters"]["failed"] == 1
    assert operator.get_status()["stop_reason"] == "MANUAL_STOP"


def test_test_discovery_bypasses_cache_and_processes_account(tmp_path):
    database = SimpleNamespace(close=lambda: None)
    calls = []
    processed = []

    def discover(**kwargs):
        calls.append(kwargs)
        return {
            "candidates": [
                {
                    "company_id": 44,
                    "company_name": "Real Precision Ltd",
                    "data_provenance": "LIVE_SEARCH_DISCOVERED",
                }
            ]
        }

    operator = _operator(tmp_path, db_factory=lambda: database)
    operator._discovery_fn = discover
    request_counts = iter([10, 11])
    operator._serper_live_request_count = lambda: next(request_counts)
    operator._process_production_account = lambda db, account, account_key: processed.append((account, account_key))

    assert operator._run_discovery_cycle() is True
    assert calls[0]["use_cache"] is False
    assert processed[0][0]["company_name"] == "Real Precision Ltd"
    assert operator.get_status()["provider_usage"]["Serper"] == 1


def test_manual_stop_finishes_current_safe_work_and_finalizes(tmp_path):
    entered = threading.Event()
    release = threading.Event()

    class BlockingOperator(SalesoorjaOperator):
        def _process_synthetic_account(self):
            entered.set()
            release.wait(timeout=5)
            return super()._process_synthetic_account()

    base = _operator(tmp_path, SALESOORJA_MODE="SMOKE")
    operator = BlockingOperator(
        settings_obj=base.settings,
        state_path=base.state_path,
        report_dir=base.report_dir,
        now=lambda: NOW,
        db_factory=None,
        rediff_adapter=base._rediff,
    )

    operator.start_run(background=True)
    assert entered.wait(timeout=3)
    stop_result = operator.stop_run(wait=False)
    assert stop_result["stopped"] is True
    assert stop_result["status"]["status"] == "STOPPING"
    release.set()
    operator._thread.join(timeout=5)

    status = operator.get_status()
    assert status["status"] == "STOPPED"
    assert status["stop_reason"] == "MANUAL_STOP"
    assert status["counters"]["emails_sent"] == 0
    assert Path(status["report_path"]).is_file()


def test_production_mode_requires_all_explicit_safety_switches(tmp_path):
    operator = _operator(tmp_path, SALESOORJA_MODE="PRODUCTION")

    result = operator.start_run(background=False)

    assert result["started"] is False
    assert result["reason"] == "PRODUCTION_GUARD_FAILED"
    assert "REAL_OUTREACH_ENABLED=true" in result["errors"]
    assert "OUTBOUND_TEST_MODE=false" in result["errors"]
    assert operator.get_status()["counters"]["emails_sent"] == 0


def test_invalid_mode_is_rejected_without_starting(tmp_path):
    operator = _operator(tmp_path, SALESOORJA_MODE="UNSAFE")

    result = operator.start_run(background=False)

    assert result == {
        "started": False,
        "reason": "INVALID_MODE",
        "errors": ["SALESOORJA_MODE must be SMOKE, TEST, or PRODUCTION"],
    }


def test_controlled_transport_receipt_persists_and_blocks_restart_duplicate(tmp_path):
    calls = []

    def worker(command, **kwargs):
        calls.append(command)
        result_path = Path(command[command.index("--result") + 1])
        result_path.write_text(
            '{"transport_status":"SENT","message_id":null,"error":null}',
            encoding="utf-8",
        )
        return SimpleNamespace(returncode=0, stdout="", stderr="")

    system = tmp_path / "Rediff_Email_System"
    system.mkdir()
    (system / "campaign_runner.py").write_text("# marker\n", encoding="utf-8")
    bridge = RediffTransportBridge(
        system_path=system,
        run_dir=tmp_path / "transport-runs",
        enabled=True,
        process_runner=worker,
        now=lambda: NOW,
    )
    operator = _operator(tmp_path, transport_bridge=bridge)

    first = operator.execute_controlled_transport_test(authorization=TEST_TRANSPORT_AUTHORIZATION)

    assert first["handoff_status"] == "HANDOFF_CREATED"
    assert first["quality_score"] >= 85
    assert first["receipt"]["transport_status"] == "SENT"
    assert first["status"]["counters"]["emails_sent"] == 1
    assert first["status"]["counters"]["test_emails_sent"] == 1
    assert first["status"]["counters"]["production_emails_sent"] == 0
    assert first["status"]["counters"]["real_prospect_emails_sent"] == 0
    assert first["status"]["transport_receipt_count"] == 1

    restarted = _operator(tmp_path, transport_bridge=bridge)
    second = restarted.execute_controlled_transport_test(authorization=TEST_TRANSPORT_AUTHORIZATION)

    assert second["receipt"]["transport_status"] == "SUPPRESSED_DUPLICATE_TRANSPORT"
    assert second["status"]["counters"]["emails_sent"] == 0
    assert second["status"]["historical_counters"]["emails_sent"] == 1
    assert second["status"]["transport_receipt_count"] == 1
    assert len(calls) == 1
